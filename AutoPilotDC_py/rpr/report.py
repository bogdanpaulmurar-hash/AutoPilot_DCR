# -*- coding: utf-8 -*-
# ============================================================
# AutoPilot DC — RPR (Python)
# report.py — Raportul Excel al unei rulari ("dovada muncii")
#
# Un fisier .xlsx pentru fiecare rulare, cu sau fara Excel de intrare:
#   Rezultate — un rand pentru fiecare element prelucrat (rand Excel sau
#               repetare a buclei de pe primul nivel; fara bucle: un rand
#               pentru toata rutina): start, durata, status, mesaj,
#               variabilele bifate, pasii parcursi, filmul si captura erorii.
#   Pasi      — fiecare pas executat (definitia / regula), cu rezultat.
#   Rutina    — rezumatul rularii si definitia completa a rutinei.
#
# Setarile stau in routine.settings.extra["report"], deci formatul
# .rpr.json ramane compatibil (campurile necunoscute sunt ignorate).
# Raportul nu are voie sa opreasca rutina: orice eroare de scriere
# ajunge doar in log.
# ============================================================

from __future__ import annotations

import datetime as _dt
import os
import re
import threading
import time
from types import SimpleNamespace
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Tuple

from .model import Routine, RprNode, sanitize_name

REPORT_KEY = "report"

STEPS_SHEET_LIMIT = 100_000   # protectie: rulari de ore cu foarte multi pasi
CHECKPOINT_MIN_S = 5.0        # salvarea din mers: cel mult o data la 5 s


_I18N_READY = False


def _t(text: str) -> str:
    """Textele din raport / documentatie in limba aleasa in Setari (i18n.py al aplicatiei, daca exista)."""
    global _I18N_READY
    try:
        import i18n
        if not _I18N_READY:
            _I18N_READY = True
            lang = i18n.saved_language()
            if lang != i18n.LANG:
                i18n.load(lang)
        return i18n.tr(text)
    except Exception:  # noqa: BLE001 — fara i18n: romana
        return text


@dataclass
class ReportSettings:
    enabled: bool = False
    folder: str = ""                     # gol = <folderul de rezultate>\rapoarte
    variables: List[str] = field(default_factory=list)   # coloanele cu valori (nume de variabile sau {sabloane})
    steps_column: bool = True            # coloana "Pasi parcursi" in Rezultate
    steps_sheet: bool = True             # foaia "Pasi"
    definition_sheet: bool = True        # foaia "Rutina" cu definitia pasilor

    @classmethod
    def of(cls, routine: Routine) -> "ReportSettings":
        d = (routine.settings.extra or {}).get(REPORT_KEY) or {}
        s = cls()
        if isinstance(d, dict):
            for k in ("enabled", "steps_column", "steps_sheet", "definition_sheet"):
                if k in d:
                    setattr(s, k, bool(d[k]))
            s.folder = str(d.get("folder") or "")
            s.variables = [str(v) for v in (d.get("variables") or []) if str(v).strip()]
        return s

    def save_to(self, routine: Routine) -> None:
        routine.settings.extra[REPORT_KEY] = {
            "enabled": self.enabled, "folder": self.folder, "variables": list(self.variables),
            "steps_column": self.steps_column, "steps_sheet": self.steps_sheet,
            "definition_sheet": self.definition_sheet,
        }


VIDEO_FOLDER_KEY = "video_folder"   # routine.settings.extra: folderul filmelor rutinei (gol = Setari → Video)


def video_folder(routine: Routine) -> str:
    return str((routine.settings.extra or {}).get(VIDEO_FOLDER_KEY) or "").strip().strip('"')


def set_video_folder(routine: Routine, folder: str) -> None:
    folder = (folder or "").strip().strip('"')
    if folder:
        routine.settings.extra[VIDEO_FOLDER_KEY] = folder
    else:
        routine.settings.extra.pop(VIDEO_FOLDER_KEY, None)


def onedrive_folder() -> str:
    """Folderul OneDrive sincronizat (de firma, apoi personal); gol daca nu exista.
    Fisierele puse aici ajung singure in cloud prin clientul OneDrive."""
    for var in ("OneDriveCommercial", "OneDrive", "OneDriveConsumer"):
        p = os.environ.get(var, "")
        if p and os.path.isdir(p):
            return p
    return ""


def toggle_variable(routine: Routine, name: str, on: bool) -> None:
    """Bifa "in raport" de pe un pas: adauga / scoate variabila din coloanele raportului."""
    name = (name or "").strip().strip("{}")
    if not name:
        return
    s = ReportSettings.of(routine)
    present = [v for v in s.variables if v.strip("{}").lower() == name.lower()]
    if on and not present:
        s.variables.append(name)
        if not s.enabled:
            s.enabled = True
    elif not on and present:
        s.variables = [v for v in s.variables if v not in present]
    s.save_to(routine)


def has_variable(routine: Routine, name: str) -> bool:
    name = (name or "").strip().strip("{}").lower()
    return bool(name) and any(v.strip("{}").lower() == name for v in ReportSettings.of(routine).variables)


@dataclass
class _Step:
    start: _dt.datetime
    item: int
    depth: int
    title: str
    kind: str
    ok: bool
    ms: float
    message: str = ""
    screenshot: str = ""


@dataclass
class _Item:
    no: int
    label: str
    start: _dt.datetime
    end: Optional[_dt.datetime] = None
    status: str = ""
    message: str = ""
    values: Dict[str, str] = field(default_factory=dict)
    steps: List[str] = field(default_factory=list)
    video: str = ""
    video_at: float = 0.0
    screenshot: str = ""
    errors_at_start: int = 0


def _fmt_dur(seconds: float) -> str:
    seconds = max(0, int(round(seconds)))
    h, rest = divmod(seconds, 3600)
    m, s = divmod(rest, 60)
    return f"{h}:{m:02d}:{s:02d}"


def _fmt_at(seconds: float) -> str:
    m, s = divmod(max(0, int(seconds)), 60)
    return f"{m:02d}:{s:02d}"


def _column_title(spec: str) -> str:
    """"Rand.ID" -> "Rand.ID"; "{Rand.ID} / {Status}" ramane asa (fara acolade daca e o singura variabila)."""
    s = spec.strip()
    if re.fullmatch(r"\{[^{}]+\}", s):
        return s[1:-1]
    return s


class RunReport:
    """Aduna ce s-a intamplat in rulare si scrie la final fisierul Excel."""

    def __init__(self, routine: Routine, settings: ReportSettings, output_folder: str):
        self.routine = routine
        self.s = settings
        self.folder = settings.folder.strip().strip('"') or os.path.join(output_folder, "rapoarte")
        self.start = _dt.datetime.now()
        self.items: List[_Item] = []
        self.steps: List[_Step] = []
        self.depth = 0                     # >0 = suntem in interiorul unui element (rand / repetare)
        self.path = ""
        self.error = ""
        self._fixed_path = ""          # aceeasi cale pentru salvarile din mers si pentru final
        self._last_cp = 0.0
        self._cp_cost = 0.0            # cat a durat ultima salvare (rulari mari = salvari mai rare)
        self.checkpoint_error = ""
        self._cur: Optional[_Item] = None
        self._lock = threading.Lock()
        self._steps_dropped = 0

    @classmethod
    def create(cls, routine: Routine, output_folder: str) -> Optional["RunReport"]:
        s = ReportSettings.of(routine)
        return cls(routine, s, output_folder) if s.enabled else None

    # ── elemente (rand Excel / repetare) ──
    def item_started(self, label: str, errors_so_far: int, video: Tuple[str, float] = ("", 0.0)) -> None:
        with self._lock:
            it = _Item(no=len(self.items) + 1, label=label, start=_dt.datetime.now(),
                       errors_at_start=errors_so_far, video=video[0], video_at=video[1])
            self.items.append(it)
            self._cur = it
            self.depth += 1

    def item_finished(self, status: str, values: Dict[str, str], video: Tuple[str, float] = ("", 0.0)) -> None:
        with self._lock:
            self.depth = max(0, self.depth - 1)
            it = self._cur
            if it is None:
                return
            it.end = _dt.datetime.now()
            it.status, _, msg = status.partition(":")
            it.status, it.message = it.status.strip(), msg.strip()
            it.values = dict(values)
            if not it.video and video[0]:
                it.video, it.video_at = video
            self._cur = None

    def evaluate_values(self, resolve) -> Dict[str, str]:
        """Valorile coloanelor alese (resolve = VariableStore.resolve)."""
        out = {}
        for spec in self.s.variables:
            expr = spec if "{" in spec else "{" + spec.strip() + "}"
            try:
                out[spec] = str(resolve(expr))
            except Exception as ex:  # noqa: BLE001
                out[spec] = f"(eroare: {ex})"
        return out

    # ── pasi ──
    def step(self, node: RprNode, ok: bool, started: _dt.datetime, ms: float, message: str = "",
             screenshot: str = "", depth: int = 0) -> None:
        with self._lock:
            item = self._cur.no if self._cur is not None else 0
            if self._cur is not None and not node.is_container:
                self._cur.steps.append(("" if ok else "✖ ") + node.title)
                if screenshot and not self._cur.screenshot:
                    self._cur.screenshot = screenshot
            if len(self.steps) >= STEPS_SHEET_LIMIT:
                self._steps_dropped += 1
                return
            self.steps.append(_Step(started, item, depth, node.title, node.KIND_LABEL, ok, ms, message, screenshot))

    # ── final ──
    def finish(self, result, final_values: Dict[str, str], video: Tuple[str, float] = ("", 0.0)) -> str:
        """Scrie fisierul; intoarce calea (sau "" + self.error)."""
        end = _dt.datetime.now()
        with self._lock:
            if self._cur is not None:            # oprire in mijlocul unui element
                self._cur.end = end
                self._cur.status = "OPRIT" if getattr(result, "cancelled", False) else "EROARE"
                self._cur.message = getattr(result, "message", "")
                self._cur.values = dict(final_values)
                self._cur = None
            if not self.items:                   # fara bucle: un rand pentru toata rutina
                it = _Item(no=1, label=self.routine.name, start=self.start, end=end,
                           status=("OK" if result.success and not result.errors else
                                   "OPRIT" if result.cancelled else "EROARE"),
                           message=result.message, values=dict(final_values), video=video[0], video_at=0.0)
                it.steps = [("" if s.ok else "✖ ") + s.title for s in sorted(self.steps, key=lambda x: x.start)
                            if s.kind not in _CONTAINER_LABELS]
                it.screenshot = next((s.screenshot for s in self.steps if s.screenshot), "")
                self.items.append(it)
        try:
            self.path = self._write(result, end, video)
        except PermissionError:
            # fisierul salvat din mers e deschis in Excel: raportul final primeste alt nume
            try:
                self._fixed_path = ""
                self._fixed_path = self._target_path(suffix="_final")
                self.path = self._write(result, end, video)
            except Exception as ex:  # noqa: BLE001
                self.error = f"Raportul nu a putut fi scris: {ex}"
                self.path = ""
        except Exception as ex:  # noqa: BLE001 — raportul nu are voie sa opreasca rutina
            self.error = f"Raportul nu a putut fi scris: {ex}"
            self.path = ""
        return self.path

    def checkpoint(self, result, force: bool = False) -> bool:
        """Salvare din mers (dupa fiecare element): daca aplicatia ingheata sau PC-ul se opreste, raportul
        ramane cu tot ce s-a facut pana atunci. Cel mult o data la cateva secunde; la rulari mari mai rar."""
        now = time.monotonic()
        if not force and now - self._last_cp < max(CHECKPOINT_MIN_S, 3 * self._cp_cost):
            return False
        partial = SimpleNamespace(message=_t("Rulare in curs — salvare automata"),
                                  steps_executed=getattr(result, "steps_executed", 0),
                                  errors=getattr(result, "errors", 0))
        try:
            self.path = self._write(partial, _dt.datetime.now(), ("", 0.0))
            self.checkpoint_error = ""
        except Exception as ex:  # noqa: BLE001 — ex. fisierul e deschis in Excel: incercam la urmatorul element
            self.checkpoint_error = str(ex)
            return False
        finally:
            self._last_cp = time.monotonic()
            self._cp_cost = self._last_cp - now
        return True

    def _target_path(self, suffix: str = "") -> str:
        if self._fixed_path:
            return self._fixed_path
        os.makedirs(self.folder, exist_ok=True)
        name = sanitize_name(self.routine.name) or "Rutina"
        base = os.path.join(self.folder, f"Raport_{name}_{self.start:%Y%m%d_%H%M%S}{suffix}")
        path, n = base + ".xlsx", 2
        while os.path.exists(path):   # aceeasi rutina de doua ori in aceeasi secunda (pregatiri)
            path, n = f"{base}_{n}.xlsx", n + 1
        self._fixed_path = path
        return path

    def _write(self, result, end: _dt.datetime, video: Tuple[str, float]) -> str:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        bold = Font(bold=True)
        head_fill = PatternFill("solid", fgColor="DDEBF7")
        fills = {"OK": "C6EFCE", "EROARE": "FFC7CE", "SARIT": "FFEB9C", "OPRIT": "FFEB9C"}
        link_font = Font(color="0563C1", underline="single")
        wrap = Alignment(wrap_text=True, vertical="top")

        def header(ws, titles, widths):
            ws.append(titles)
            for i, w in enumerate(widths, 1):
                c = ws.cell(row=1, column=i)
                c.font, c.fill = bold, head_fill
                ws.column_dimensions[get_column_letter(i)].width = w
            ws.freeze_panes = "A2"

        def link(cell, path, text=None):
            if not path:
                return
            cell.value = text or os.path.basename(path)
            cell.hyperlink = "file:///" + os.path.abspath(path).replace("\\", "/")
            cell.font = link_font

        # 1. Rezultate
        ws = wb.active
        ws.title = _t("Rezultate")
        cols = [_t(c) for c in ["#", "Element", "Start", "Sfarsit", "Durata", "Status", "Mesaj"]]
        widths = [5, 28, 19, 19, 10, 10, 40]
        var_titles = [_column_title(v) for v in self.s.variables]
        cols += var_titles
        widths += [18] * len(var_titles)
        if self.s.steps_column:
            cols.append(_t("Pasi parcursi"))
            widths.append(50)
        cols += [_t(c) for c in ["Film", "Minut in film", "Captura eroare"]]
        widths += [30, 12, 30]
        header(ws, cols, widths)
        for it in self.items:
            dur = ((it.end or end) - it.start).total_seconds()
            row = [it.no, it.label, it.start.strftime("%d.%m.%Y %H:%M:%S"),
                   (it.end or end).strftime("%d.%m.%Y %H:%M:%S"), _fmt_dur(dur), it.status, it.message]
            row += [it.values.get(v, "") for v in self.s.variables]
            if self.s.steps_column:
                steps = it.steps if len(it.steps) <= 200 else it.steps[:200] + [f"... (+{len(it.steps) - 200})"]
                row.append("\n".join(steps)[:32000])
            row += ["", _fmt_at(it.video_at) if it.video else "", ""]
            ws.append(row)
            r = ws.max_row
            color = fills.get(it.status.upper().split()[0] if it.status else "", None)
            if color:
                ws.cell(row=r, column=6).fill = PatternFill("solid", fgColor=color)
            n = len(cols)
            link(ws.cell(row=r, column=n - 2), it.video)
            link(ws.cell(row=r, column=n), it.screenshot)
            if self.s.steps_column:
                ws.cell(row=r, column=n - 3).alignment = wrap
        ws.auto_filter.ref = ws.dimensions

        # 2. Pasi
        if self.s.steps_sheet:
            ps = wb.create_sheet(_t("Pasi"))
            header(ps, [_t(c) for c in ["Ora", "Element", "Pas (definitie / regula)", "Tip", "Rezultat", "Durata (ms)",
                                        "Mesaj", "Captura"]], [10, 9, 60, 22, 9, 11, 40, 30])
            for st in sorted(self.steps, key=lambda x: x.start):
                ps.append([st.start.strftime("%H:%M:%S"), st.item or "", "    " * st.depth + st.title, st.kind,
                           "OK" if st.ok else "EROARE", int(st.ms), st.message, ""])
                if not st.ok:
                    ps.cell(row=ps.max_row, column=5).fill = PatternFill("solid", fgColor=fills["EROARE"])
                link(ps.cell(row=ps.max_row, column=8), st.screenshot)
            if self._steps_dropped:
                ps.append(["", "", f"... inca {self._steps_dropped} pasi nescrisi (limita {STEPS_SHEET_LIMIT})"])
            ps.auto_filter.ref = ps.dimensions

        # 3. Rutina: rezumat + definitie
        rs = wb.create_sheet(_t("Rutina"))
        rs.column_dimensions["A"].width = 26
        rs.column_dimensions["B"].width = 90
        ok = sum(1 for i in self.items if i.status.upper().startswith("OK"))
        summary = [
            ("Rutina", self.routine.name),
            ("Descriere", self.routine.description),
            ("Start", self.start.strftime("%d.%m.%Y %H:%M:%S")),
            ("Sfarsit", end.strftime("%d.%m.%Y %H:%M:%S")),
            ("Durata", _fmt_dur((end - self.start).total_seconds())),
            ("Rezultat", result.message),
            ("Elemente", f"{len(self.items)} (OK: {ok}, altele: {len(self.items) - ok})"),
            ("Pasi executati", str(result.steps_executed)),
            ("Erori", str(result.errors)),
            ("Utilizator / PC", f"{os.environ.get('USERNAME') or os.environ.get('USER') or ''} / "
                                f"{os.environ.get('COMPUTERNAME') or ''}".strip(" /")),
        ]
        if video[0]:
            summary.append(("Film", video[0]))
        for k, v in summary:
            rs.append([_t(k), v])
            rs.cell(row=rs.max_row, column=1).font = bold
        if video[0]:
            link(rs.cell(row=rs.max_row, column=2), video[0], video[0])
        if self.routine.variables:
            rs.append([])
            rs.append([_t("Variabile de start")])
            rs.cell(row=rs.max_row, column=1).font = bold
            for k, v in self.routine.variables.items():
                rs.append([k, v])
        if self.s.definition_sheet:
            rs.append([])
            rs.append([_t("Definitia rutinei")])
            rs.cell(row=rs.max_row, column=1).font = bold
            for depth, node in _walk(self.routine.steps):
                prefix = "" if node.enabled else "(dezactivat) "
                rs.append([_t(node.KIND_LABEL), "    " * depth + prefix + node.title +
                           (f"   // {node.comment}" if node.comment.strip() else "")])

        path = self._target_path()
        tmp = path + ".tmp"
        wb.save(tmp)
        os.replace(tmp, path)
        return path


def _walk(nodes, depth=0):
    for n in nodes:
        yield depth, n
        for name, lst in n.branches():
            if len(n.branches()) > 1:
                yield depth + 1, _Branch(name)
            yield from _walk(lst, depth + (2 if len(n.branches()) > 1 else 1))


class _Branch:
    """Eticheta de ramura in definitie (ex. "Altfel")."""
    KIND_LABEL = ""
    enabled = True
    comment = ""

    def __init__(self, name):
        self.title = f"[{name}]"


_CONTAINER_LABELS = set()


def _init_container_labels():
    from .model import NODE_TYPES
    for cls in NODE_TYPES.values():
        if cls._child_lists:
            _CONTAINER_LABELS.add(cls.KIND_LABEL)


_init_container_labels()
