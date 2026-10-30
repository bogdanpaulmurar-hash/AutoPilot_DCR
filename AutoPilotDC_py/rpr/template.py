# -*- coding: utf-8 -*-
# ============================================================
# AutoPilot DC v4.23 — RPR (Python)
# template.py — 📤 Exporta sablon Excel pentru o rutina cu "Pentru fiecare rand Excel"
#
# Sablonul are exact coloanele pe care le asteapta rutina:
#   Id (coloana cheie, formatata ca text: zerourile din fata si Id-urile lungi raman cum le scrii)
#   + coloanele in care scrie rutina (pasii "Scrie in Excel" din bucla)
#   + coloana de status.
# Foaia "Instructiuni" explica ce se completeaza; foaia ascunsa "_autopilot" tine rutina si versiunea.
# ============================================================

from __future__ import annotations

import datetime as _dt
import re
from typing import List, Optional

from .drivers.base import RprStepException
from .excel import ID_ALIASES, _compact, _require_openpyxl
from .lang import ft
from .model import ForEachRowNode, Routine, RprNode, WriteExcelNode

try:
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    openpyxl = None

DEFAULT_KEY = "Id"
DATA_ROWS = 2000                 # cate randuri primesc formatul text pe coloana Id
META_SHEET = "_autopilot"
_FILL_KEY, _FILL_OUT, _FILL_STATUS = "BDD7EE", "C6EFCE", "D9D9D9"


def _is_letter(col: str) -> bool:
    """"D" = litera de coloana; "Id", "Nr", "Cod" sunt titluri."""
    col = (col or "").strip()
    return bool(re.fullmatch(r"[A-Za-z]{1,3}", col)) and _compact(col) not in ID_ALIASES


def first_loop(routine: Routine) -> Optional[ForEachRowNode]:
    """Prima bucla Excel activa (inclusiv in grupuri / Daca / Incearca)."""
    def walk(nodes: List[RprNode]):
        for n in nodes:
            if isinstance(n, ForEachRowNode) and n.enabled:
                return n
            for _, lst in n.branches():
                found = walk(lst)
                if found is not None:
                    return found
        return None
    return walk(routine.steps)


def _written_columns(nodes: List[RprNode], out: List[str]) -> None:
    for n in nodes:
        if isinstance(n, ForEachRowNode):   # o bucla in bucla scrie in alt fisier
            continue
        if isinstance(n, WriteExcelNode) and n.enabled:
            c = (n.column or "").strip()
            if c and "{" not in c and not _is_letter(c) and c.lower() not in (x.lower() for x in out):
                out.append(c)
        for _, lst in n.branches():
            _written_columns(lst, out)


def template_columns(f: ForEachRowNode) -> tuple:
    """(cheia, coloanele scrise de rutina, coloana de status) — in ordinea din sablon."""
    key = f.key_column.strip()
    if not key or _is_letter(key) or "{" in key:
        key = DEFAULT_KEY
    status = f.status_column.strip()
    written: List[str] = []
    _written_columns(f.body, written)
    written = [c for c in written if c.lower() not in (key.lower(), status.lower())]
    return key, written, status


def export_template(routine: Routine, path: str, app_version: str = "") -> str:
    """Scrie sablonul la `path` si intoarce calea. RprStepException cu mesaj clar daca rutina nu are bucla Excel."""
    _require_openpyxl()
    f = first_loop(routine)
    if f is None:
        raise RprStepException("Rutina nu are pasul \"Pentru fiecare rand Excel\": nu stiu ce coloane sa pun in sablon.")
    if f.header_row <= 0:
        raise RprStepException("Bucla Excel citeste fara rand de titluri (Randul cu titlurile = 0). "
                               "Pune 1 ca sa poti folosi sablonul.")
    key, written, status = template_columns(f)
    hr = f.header_row
    sheet = f.sheet.strip()
    if not sheet or "{" in sheet or len(sheet) > 31 or re.search(r"[\[\]:*?/\\]", sheet):
        sheet = ft("Date")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet

    cols = [(key, _FILL_KEY, ft("Completezi tu: Id-ul dupa care rutina cauta. Orice format (cifre, litere, "
                                "zerouri in fata); coloana e text, Excel nu il schimba."))]
    cols += [(c, _FILL_OUT, ft("Completeaza rutina (pasul \"Scrie in Excel\"). Lasa gol.")) for c in written]
    if status:
        cols.append((status, _FILL_STATUS, ft("Completeaza rutina: OK / EROARE / SARIT. Randurile OK se sar la "
                                              "reluare; sterge statusul ca sa rulezi randul din nou.")))

    for i, (title, fill, note) in enumerate(cols, start=1):
        cell = ws.cell(row=hr, column=i, value=title)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(fill_type="solid", start_color=fill, end_color=fill)
        cell.alignment = Alignment(vertical="center")
        cell.comment = Comment(note, "AutoPilot DC", width=260, height=110)
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(40, len(title) + 6))
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width, 22)

    # Id ca text: 00123 ramane 00123, 123456789012345678 nu devine 1,23E+17
    ws.column_dimensions["A"].number_format = "@"
    for r in range(hr + 1, hr + 1 + DATA_ROWS):
        ws.cell(row=r, column=1).number_format = "@"
    # avertisment (nu blocare) la Id repetat, chiar in timp ce scrii
    rng = f"A{hr + 1}:A{hr + DATA_ROWS}"
    dv = DataValidation(type="custom", formula1=f"COUNTIF($A${hr + 1}:$A${hr + DATA_ROWS},A{hr + 1})<=1",
                        allow_blank=True, showErrorMessage=True, errorStyle="warning",
                        errorTitle=ft("Id repetat"), error=ft("Acest Id exista deja in lista."))
    dv.add(rng)
    ws.add_data_validation(dv)
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)

    # Instructiuni
    wi = wb.create_sheet(ft("Instructiuni"))
    lines = [
        ft("Sablon pentru rutina: {0}").replace("{0}", routine.name),
        "",
        ft("1. Pe foaia \"{0}\" completezi doar coloana {1}, cate un Id pe rand, fara randuri goale intre ele.")
        .replace("{0}", sheet).replace("{1}", key),
        ft("2. Coloanele verzi le completeaza rutina. Coloana gri e statusul randului."),
        ft("3. In Studio, la \"Pentru fiecare rand Excel\", alegi acest fisier si rulezi."),
        ft("4. Inainte de primul click, AutoPilot verifica randurile fara Id si Id-urile repetate si te intreaba."),
        "",
        ft("Coloana Id poate avea si alt titlu: ID, UEID, Cod, Nr — AutoPilot o recunoaste."),
    ]
    for i, t in enumerate(lines, start=1):
        wi.cell(row=i, column=1, value=t)
    wi["A1"].font = Font(bold=True, size=12)
    wi.column_dimensions["A"].width = 110

    # Date pentru aplicatie (ascunse)
    wm = wb.create_sheet(META_SHEET)
    meta = [("routine", routine.name), ("app_version", app_version),
            ("exported", _dt.datetime.now().strftime("%Y-%m-%d %H:%M")), ("sheet", sheet), ("header_row", hr),
            ("key_column", key), ("output_columns", " | ".join(written)), ("status_column", status)]
    for i, (k, v) in enumerate(meta, start=1):
        wm.cell(row=i, column=1, value=k)
        wm.cell(row=i, column=2, value=v)
    wm.sheet_state = "hidden"

    wb.active = 0
    try:
        wb.save(path)
    except PermissionError as ex:
        raise RprStepException(f"Nu pot salva sablonul: fisierul e deschis in alt program. Inchide-l in Excel. ({ex})")
    return path
