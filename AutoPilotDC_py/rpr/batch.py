# -*- coding: utf-8 -*-
# ============================================================
# AutoPilot DC — RPR (Python)
# batch.py — Managerul de rutine si "Pregatirile"
#
#   RoutineLibrary   — rutinele (.rpr.json) dintr-un folder ales, citite la
#                      pornire (si la cerere), cu nume, descriere, pasi.
#   Preparation      — o lista de rutine rulate una dupa alta (ex. un sandbox:
#                      creeaza useri, apoi roluri, apoi alte obiecte), salvata
#                      ca <nume>.pregatire.json in acelasi folder.
#   PreparationRunner— ruleaza lista; pentru fiecare rutina: raportul Excel,
#                      filmul, apoi la final un folder local cu:
#                        Pregatire_<nume>.xlsx  (Rezumat + Definitii)
#                        Documentatie.html      (tot ce s-a facut, in detaliu)
#                        rapoartele fiecarei rutine
#
# Rularea efectiva a unei rutine e data din afara (run_one), ca Studio sa
# pastreze intrebarile si sistemul de preventie, iar rularea programata sa
# mearga fara interfata.
# ============================================================

from __future__ import annotations

import copy
import datetime as _dt
import html
import json
import os
import threading
from dataclasses import dataclass, field
from typing import Callable, Dict, List, Optional, Tuple

from .model import Routine, sanitize_name
from .report import REPORT_KEY, ReportSettings, _fmt_dur, _t, _walk

PREP_EXT = ".pregatire.json"
ROUTINE_EXT = ".rpr.json"


# ─────────────────────────────────────────────────────────────
# 1. Biblioteca de rutine (folder)
# ─────────────────────────────────────────────────────────────
@dataclass
class RoutineEntry:
    path: str
    name: str
    description: str = ""
    steps: int = 0
    modified: _dt.datetime = field(default_factory=_dt.datetime.now)
    error: str = ""

    @property
    def folder(self) -> str:
        return os.path.dirname(self.path)


class RoutineLibrary:
    """Rutinele dintr-un folder (si subfoldere). Citirea e rapida si nu opreste aplicatia la un fisier stricat."""

    def __init__(self, folder: str = ""):
        self.folder = folder
        self.routines: List[RoutineEntry] = []
        self.preparations: List[str] = []
        self.scanned_at: Optional[_dt.datetime] = None
        self._lock = threading.Lock()

    def scan(self, folder: Optional[str] = None, max_depth: int = 4) -> "RoutineLibrary":
        if folder is not None:
            self.folder = folder
        routines, preps = [], []
        root = (self.folder or "").strip().strip('"')
        if root and os.path.isdir(root):
            base_depth = root.rstrip("\\/").count(os.sep)
            for dirpath, dirnames, filenames in os.walk(root):
                if dirpath.count(os.sep) - base_depth >= max_depth:
                    dirnames[:] = []
                dirnames[:] = sorted(d for d in dirnames if not d.startswith((".", "_")))
                for fn in sorted(filenames, key=str.lower):
                    p = os.path.join(dirpath, fn)
                    low = fn.lower()
                    if low.endswith(PREP_EXT):
                        preps.append(p)
                    elif low.endswith(ROUTINE_EXT):
                        routines.append(_read_entry(p))
        with self._lock:
            self.routines, self.preparations = routines, preps
            self.scanned_at = _dt.datetime.now()
        return self

    def find(self, path: str) -> Optional[RoutineEntry]:
        key = os.path.normcase(os.path.abspath(path))
        return next((r for r in self.routines if os.path.normcase(os.path.abspath(r.path)) == key), None)


def _read_entry(path: str) -> RoutineEntry:
    name = os.path.basename(path)[:-len(ROUTINE_EXT)]
    try:
        with open(path, encoding="utf-8-sig") as f:
            d = json.load(f)
        steps = 0

        def count(nodes):
            nonlocal steps
            for n in nodes or []:
                if isinstance(n, dict):
                    steps += 1
                    for k in ("body", "then", "else", "catch"):
                        count(n.get(k))
        count(d.get("steps") or d.get("Steps"))
        return RoutineEntry(path, str(d.get("name") or d.get("Name") or name),
                            str(d.get("description") or d.get("Description") or ""), steps,
                            _dt.datetime.fromtimestamp(os.path.getmtime(path)))
    except Exception as ex:  # noqa: BLE001
        return RoutineEntry(path, name, error=str(ex))


# ─────────────────────────────────────────────────────────────
# 2. Pregatirea (lista de rutine)
# ─────────────────────────────────────────────────────────────
@dataclass
class PrepItem:
    path: str
    enabled: bool = True
    stop_on_error: bool = True                  # rutina esuata opreste pregatirea (debifat = trece la urmatoarea)
    variables: Dict[str, str] = field(default_factory=dict)   # suprascrie variabilele de start ale rutinei

    @property
    def name(self) -> str:
        n = os.path.basename(self.path)
        return n[:-len(ROUTINE_EXT)] if n.lower().endswith(ROUTINE_EXT) else n


@dataclass
class Preparation:
    name: str = "Pregatire noua"
    description: str = ""
    items: List[PrepItem] = field(default_factory=list)
    output_folder: str = ""                      # gol = <rezultate>\pregatiri
    video: bool = True                           # film pentru fiecare rutina
    report: bool = True                          # raport Excel pentru fiecare rutina (si cand rutina nu il are pornit)
    path: str = ""                               # fisierul .pregatire.json (nu se salveaza in el)

    def to_dict(self) -> dict:
        return {"format": 1, "name": self.name, "description": self.description,
                "output_folder": self.output_folder, "video": self.video, "report": self.report,
                "items": [{"path": i.path, "enabled": i.enabled, "stop_on_error": i.stop_on_error,
                           "variables": dict(i.variables)} for i in self.items]}

    @classmethod
    def from_dict(cls, d: dict, path: str = "") -> "Preparation":
        p = cls(name=str(d.get("name") or "Pregatire"), description=str(d.get("description") or ""),
                output_folder=str(d.get("output_folder") or ""), video=bool(d.get("video", True)),
                report=bool(d.get("report", True)), path=path)
        for it in d.get("items") or []:
            if isinstance(it, dict) and it.get("path"):
                p.items.append(PrepItem(str(it["path"]), bool(it.get("enabled", True)),
                                        bool(it.get("stop_on_error", True)),
                                        {str(k): str(v) for k, v in (it.get("variables") or {}).items()}))
        return p

    @classmethod
    def load(cls, path: str) -> "Preparation":
        with open(path, encoding="utf-8-sig") as f:
            return cls.from_dict(json.load(f), path)

    def save(self, path: Optional[str] = None) -> str:
        path = path or self.path
        if not path:
            raise ValueError("Pregatirea nu are inca un fisier.")
        os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
        tmp = path + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(self.to_dict(), f, ensure_ascii=False, indent=2)
        os.replace(tmp, path)
        self.path = path
        return path

    def resolve(self, item: PrepItem) -> str:
        """Caile relative sunt fata de fisierul pregatirii (folderul poate fi mutat / sincronizat in Drive)."""
        p = item.path.strip().strip('"')
        if not os.path.isabs(p) and self.path:
            p = os.path.join(os.path.dirname(os.path.abspath(self.path)), p)
        return os.path.normpath(p)

    def relative(self, routine_path: str) -> str:
        """Calea salvata in fisier: relativa daca rutina e langa pregatire (sau in subfoldere)."""
        if self.path:
            base = os.path.dirname(os.path.abspath(self.path))
            try:
                rel = os.path.relpath(os.path.abspath(routine_path), base)
                if not rel.startswith(".."):
                    return rel
            except ValueError:   # alt disc
                pass
        return os.path.abspath(routine_path)


# ─────────────────────────────────────────────────────────────
# 3. Rularea
# ─────────────────────────────────────────────────────────────
@dataclass
class PrepStepResult:
    item: PrepItem
    name: str
    path: str
    start: _dt.datetime
    end: Optional[_dt.datetime] = None
    status: str = ""             # OK / EROARE / OPRIT / SARIT / LIPSA
    message: str = ""
    steps: int = 0
    errors: int = 0
    rows_ok: int = 0
    rows_failed: int = 0
    report: str = ""
    video: str = ""
    routine: Optional[Routine] = None


@dataclass
class PrepResult:
    name: str
    start: _dt.datetime
    end: Optional[_dt.datetime] = None
    folder: str = ""
    summary: str = ""
    documentation: str = ""
    steps: List[PrepStepResult] = field(default_factory=list)
    cancelled: bool = False

    @property
    def success(self) -> bool:
        return not self.cancelled and all(s.status in ("OK", "SARIT") for s in self.steps)

    @property
    def message(self) -> str:
        ok = sum(1 for s in self.steps if s.status == "OK")
        failed = sum(1 for s in self.steps if s.status in ("EROARE", "LIPSA"))
        text = f"Pregatirea \"{self.name}\": {ok} rutine OK, {failed} cu erori, din {len(self.steps)}."
        return ("Oprita de utilizator. " + text) if self.cancelled else text


# run_one(routine, item, run_folder) -> (RunResult, cale film)
RunOne = Callable[[Routine, PrepItem, str], Tuple[object, str]]


class PreparationRunner:
    def __init__(self, prep: Preparation, run_one: RunOne, output_root: str,
                 load_routine: Optional[Callable[[str], Routine]] = None,
                 on_progress: Optional[Callable[[int, int, str], None]] = None):
        self.prep = prep
        self._run_one = run_one
        self._root = prep.output_folder.strip().strip('"') or os.path.join(output_root, "pregatiri")
        self._load = load_routine or _load_routine
        self._progress = on_progress or (lambda i, n, name: None)
        self._cancel = threading.Event()

    def cancel(self) -> None:
        self._cancel.set()

    @property
    def cancelled(self) -> bool:
        return self._cancel.is_set()

    def run(self) -> PrepResult:
        start = _dt.datetime.now()
        folder = os.path.join(self._root, f"{sanitize_name(self.prep.name) or 'Pregatire'}_{start:%Y%m%d_%H%M%S}")
        os.makedirs(folder, exist_ok=True)
        res = PrepResult(self.prep.name, start, folder=folder)
        items = [i for i in self.prep.items if i.enabled]
        stop_reason = ""
        for n, item in enumerate(items, 1):
            path = self.prep.resolve(item)
            step = PrepStepResult(item, item.name, path, _dt.datetime.now())
            res.steps.append(step)
            if stop_reason or self._cancel.is_set():
                step.status, step.message, step.end = "SARIT", stop_reason or "Pregatirea a fost oprita.", step.start
                continue
            self._progress(n, len(items), item.name)
            try:
                routine = self._load(path)
            except Exception as ex:  # noqa: BLE001
                step.status, step.message, step.end = "LIPSA", f"Rutina nu poate fi citita: {ex}", _dt.datetime.now()
                if item.stop_on_error:
                    stop_reason = f"Oprit dupa eroarea rutinei \"{item.name}\"."
                continue
            step.name = routine.name or item.name
            routine = copy.deepcopy(routine)
            for k, v in item.variables.items():
                routine.variables[k] = v
            if self.prep.report:
                rs = ReportSettings.of(routine)
                rs.enabled = True
                rs.folder = os.path.join(folder, "rapoarte")
                rs.save_to(routine)
            elif (routine.settings.extra or {}).get(REPORT_KEY):
                rs = ReportSettings.of(routine)
                if rs.enabled and not rs.folder:
                    rs.folder = os.path.join(folder, "rapoarte")
                    rs.save_to(routine)
            step.routine = routine
            try:
                result, video = self._run_one(routine, item, folder)
            except Exception as ex:  # noqa: BLE001
                result, video = None, ""
                step.message = f"Eroare interna: {ex}"
            step.end = _dt.datetime.now()
            step.video = video or ""
            if result is None:
                step.status = "OPRIT" if self._cancel.is_set() else "EROARE"
            else:
                step.message = getattr(result, "message", "")
                step.steps = getattr(result, "steps_executed", 0)
                step.errors = getattr(result, "errors", 0)
                step.rows_ok = getattr(result, "rows_ok", 0)
                step.rows_failed = getattr(result, "rows_failed", 0)
                step.report = getattr(result, "report_path", "")
                if getattr(result, "cancelled", False):
                    step.status = "OPRIT"
                    self._cancel.set()
                else:
                    step.status = "OK" if getattr(result, "success", False) else "EROARE"
            if step.status == "EROARE" and item.stop_on_error:
                stop_reason = f"Oprit dupa eroarea rutinei \"{step.name}\"."
        res.cancelled = self._cancel.is_set()
        res.end = _dt.datetime.now()
        try:
            res.summary = write_summary_xlsx(res, self.prep)
        except Exception as ex:  # noqa: BLE001
            res.summary = ""
            res.documentation = ""
            _write_text(os.path.join(folder, "EROARE_rezumat.txt"), str(ex))
        try:
            res.documentation = write_documentation(res, self.prep)
        except Exception as ex:  # noqa: BLE001
            _write_text(os.path.join(folder, "EROARE_documentatie.txt"), str(ex))
        return res


def _load_routine(path: str) -> Routine:
    from .storage import RoutineStore
    return RoutineStore.load(path)


def _write_text(path: str, text: str) -> None:
    try:
        with open(path, "w", encoding="utf-8") as f:
            f.write(text)
    except OSError:
        pass


def _uri(path: str) -> str:
    return "file:///" + os.path.abspath(path).replace("\\", "/")


# ─────────────────────────────────────────────────────────────
# 4. Rezumatul Excel si documentatia
# ─────────────────────────────────────────────────────────────
def write_summary_xlsx(res: PrepResult, prep: Preparation) -> str:
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    bold, head = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
    link_font = Font(color="0563C1", underline="single")
    fills = {"OK": "C6EFCE", "EROARE": "FFC7CE", "LIPSA": "FFC7CE", "OPRIT": "FFEB9C", "SARIT": "FFEB9C"}

    ws = wb.active
    ws.title = _t("Rezumat")
    end = res.end or _dt.datetime.now()
    info = [("Pregatire", prep.name), ("Descriere", prep.description),
            ("Start", res.start.strftime("%d.%m.%Y %H:%M:%S")), ("Sfarsit", end.strftime("%d.%m.%Y %H:%M:%S")),
            ("Durata", _fmt_dur((end - res.start).total_seconds())), ("Rezultat", res.message)]
    for k, v in info:
        ws.append([_t(k), v])
        ws.cell(row=ws.max_row, column=1).font = bold
    ws.append([])
    titles = [_t(c) for c in ["#", "Rutina", "Fisier", "Start", "Sfarsit", "Durata", "Status", "Mesaj", "Pasi", "Erori",
              "Randuri OK", "Randuri esuate", "Raport", "Film"]]
    widths = [5, 30, 36, 19, 19, 10, 10, 48, 7, 7, 11, 14, 30, 30]
    ws.append(titles)
    hr = ws.max_row
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=hr, column=i)
        c.font, c.fill = bold, head
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.column_dimensions["A"].width = 12
    for n, s in enumerate(res.steps, 1):
        e = s.end or s.start
        ws.append([n, s.name, s.path, s.start.strftime("%d.%m.%Y %H:%M:%S"), e.strftime("%d.%m.%Y %H:%M:%S"),
                   _fmt_dur((e - s.start).total_seconds()), s.status, s.message, s.steps, s.errors, s.rows_ok,
                   s.rows_failed, "", ""])
        r = ws.max_row
        if s.status in fills:
            ws.cell(row=r, column=7).fill = PatternFill("solid", fgColor=fills[s.status])
        for col, p in ((13, s.report), (14, s.video)):
            if p:
                c = ws.cell(row=r, column=col)
                c.value, c.hyperlink, c.font = os.path.basename(p), _uri(p), link_font
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)

    ds = wb.create_sheet(_t("Definitii"))
    ds.column_dimensions["A"].width = 30
    ds.column_dimensions["B"].width = 24
    ds.column_dimensions["C"].width = 90
    ds.append([_t("Rutina"), _t("Tip pas"), _t("Pas (definitie / regula)")])
    for c in ds[1]:
        c.font, c.fill = bold, head
    for s in res.steps:
        if s.routine is None:
            continue
        for depth, node in _walk(s.routine.steps):
            ds.append([s.name, _t(node.KIND_LABEL), "    " * depth + ("" if node.enabled else "(dezactivat) ") + node.title])

    path = os.path.join(res.folder, f"Pregatire_{sanitize_name(prep.name) or 'Pregatire'}.xlsx")
    wb.save(path)
    return path


def write_documentation(res: PrepResult, prep: Preparation) -> str:
    e = html.escape
    end = res.end or _dt.datetime.now()
    css = ("body{font-family:Segoe UI,Arial,sans-serif;margin:24px;color:#222;max-width:1100px}"
           "h1{color:#1F4E79}h2{color:#1F4E79;border-bottom:1px solid #ccc;padding-bottom:4px;margin-top:32px}"
           "table{border-collapse:collapse;width:100%;margin:8px 0}td,th{border:1px solid #ddd;padding:4px 8px;"
           "text-align:left;vertical-align:top}th{background:#DDEBF7}.OK{background:#C6EFCE}"
           ".EROARE,.LIPSA{background:#FFC7CE}.OPRIT,.SARIT{background:#FFEB9C}"
           "pre{background:#f6f8fa;padding:8px;white-space:pre-wrap;font-family:Consolas,monospace;font-size:13px}")
    out = [f"<!doctype html><html lang='ro'><head><meta charset='utf-8'><title>{e(prep.name)}</title>"
           f"<style>{css}</style></head><body>",
           f"<h1>{e(prep.name)}</h1>"]
    if prep.description:
        out.append(f"<p>{e(prep.description)}</p>")
    out.append("<table>")
    for k, v in [("Start", res.start.strftime("%d.%m.%Y %H:%M:%S")), ("Sfarsit", end.strftime("%d.%m.%Y %H:%M:%S")),
                 ("Durata", _fmt_dur((end - res.start).total_seconds())), ("Rezultat", res.message)]:
        out.append(f"<tr><th>{e(_t(k))}</th><td>{e(v)}</td></tr>")
    if res.summary:
        out.append(f"<tr><th>{e(_t('Rezumat Excel'))}</th><td><a href='{_uri(res.summary)}'>{e(os.path.basename(res.summary))}"
                   "</a></td></tr>")
    out.append(f"</table><h2>{e(_t('Rutinele rulate'))}</h2><table><tr>" + "".join(
        f"<th>{e(_t(h))}</th>" for h in ("#", "Rutina", "Start", "Durata", "Status", "Mesaj")) + "</tr>")
    for n, s in enumerate(res.steps, 1):
        d = _fmt_dur(((s.end or s.start) - s.start).total_seconds())
        out.append(f"<tr><td>{n}</td><td><a href='#r{n}'>{e(s.name)}</a></td><td>{s.start:%H:%M:%S}</td><td>{d}</td>"
                   f"<td class='{e(s.status)}'>{e(s.status)}</td><td>{e(s.message)}</td></tr>")
    out.append("</table>")
    for n, s in enumerate(res.steps, 1):
        out.append(f"<h2 id='r{n}'>{n}. {e(s.name)}</h2>")
        if s.routine is not None and s.routine.description:
            out.append(f"<p>{e(s.routine.description)}</p>")
        rows = [("Fisier", e(s.path)), ("Status", f"<span class='{e(s.status)}'>{e(s.status)}</span> {e(s.message)}"),
                ("Start", f"{s.start:%d.%m.%Y %H:%M:%S}"),
                ("Durata", _fmt_dur(((s.end or s.start) - s.start).total_seconds())),
                ("Pasi executati / erori", f"{s.steps} / {s.errors}")]
        if s.rows_ok or s.rows_failed:
            rows.append(("Randuri Excel OK / esuate", f"{s.rows_ok} / {s.rows_failed}"))
        if s.report:
            rows.append(("Raport", f"<a href='{_uri(s.report)}'>{e(os.path.basename(s.report))}</a>"))
        if s.video:
            rows.append(("Film", f"<a href='{_uri(s.video)}'>{e(os.path.basename(s.video))}</a>"))
        if s.item.variables:
            rows.append(("Variabile date de pregatire", e(", ".join(f"{k} = {v}" for k, v in s.item.variables.items()))))
        out.append("<table>" + "".join(f"<tr><th>{e(_t(k))}</th><td>{v}</td></tr>" for k, v in rows) + "</table>")
        if s.routine is not None:
            if s.routine.variables:
                out.append(f"<p><b>{e(_t('Variabile de start'))}</b></p><table>" + "".join(
                    f"<tr><td>{e(k)}</td><td>{e(v)}</td></tr>" for k, v in s.routine.variables.items()) + "</table>")
            lines = []
            for depth, node in _walk(s.routine.steps):
                text = "    " * depth + ("(dezactivat) " if not node.enabled else "") + node.title
                if node.comment.strip():
                    text += f"   // {node.comment}"
                lines.append(text)
            out.append(f"<p><b>{e(_t('Pasii rutinei'))}</b></p><pre>" + e("\n".join(lines)) + "</pre>")
    gen = _t(f"Generat de AutoPilot DC la {end:%d.%m.%Y %H:%M:%S}.")
    out.append(f"<p style='color:#888'>{e(gen)}</p></body></html>")
    path = os.path.join(res.folder, "Documentatie.html")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(out))
    return path


# ─────────────────────────────────────────────────────────────
# 5. Biblioteca comuna aplicatiei (citita la pornire, pe un fir separat)
# ─────────────────────────────────────────────────────────────
_LIBRARY: Optional[RoutineLibrary] = None


def shared_library() -> RoutineLibrary:
    global _LIBRARY
    if _LIBRARY is None:
        _LIBRARY = RoutineLibrary()
    return _LIBRARY


def preload(folder: str, on_done: Optional[Callable[[RoutineLibrary], None]] = None) -> None:
    """Citeste folderul rutinelor in fundal (pornirea aplicatiei nu asteapta)."""
    def work():
        try:
            lib = shared_library().scan(folder)
        except Exception:  # noqa: BLE001
            return
        if on_done:
            try:
                on_done(lib)
            except Exception:  # noqa: BLE001
                pass
    threading.Thread(target=work, daemon=True, name="RoutineLibrary").start()
