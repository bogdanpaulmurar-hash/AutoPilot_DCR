# -*- coding: utf-8 -*-
# ============================================================
# AutoPilot DC v4.23 — Teste: 📤 sablon Excel, coloana Id sub alt nume, verificarea Id-urilor.
#   Rulare:  py -m unittest tests.test_template
# ============================================================

import os
import sys
import tempfile
import unittest
import uuid

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl  # noqa: E402

from rpr import (ForEachRowNode, GroupNode, IfNode, RprStepException, SetVariableNode,  # noqa: E402
                 VariableOperation, WriteExcelNode)
from rpr.engine import KEY_CHECK_BUTTONS  # noqa: E402
from rpr.excel import ExcelTable  # noqa: E402
from rpr import template  # noqa: E402
from test_rpr import create, new_routine  # noqa: E402


def _tmp():
    return os.path.join(tempfile.gettempdir(), f"apdc_tpl_{uuid.uuid4().hex}.xlsx")


def _file(rows, header=("UEID", "Nume")):
    path = _tmp()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(header))
    for r in rows:
        ws.append(list(r))
    wb.save(path)
    return path


class ExportTemplate(unittest.TestCase):
    def setUp(self):
        self.paths = []

    def tearDown(self):
        for p in self.paths:
            if os.path.exists(p):
                os.remove(p)

    def _routine(self):
        each = ForEachRowNode(file_path="", key_column="Id", status_column="Status")
        g = GroupNode()
        g.body.append(WriteExcelNode(column="Pret", value="{Pret}"))
        each.body.append(g)
        each.body.append(IfNode())
        each.body[1].then.append(WriteExcelNode(column="Mesaj", value="x"))
        each.body.append(WriteExcelNode(column="pret", value="dublura"))       # aceeasi coloana, alt caz
        each.body.append(WriteExcelNode(column="{Coloana}", value="variabila"))  # nu se stie dinainte
        each.body.append(WriteExcelNode(column="D", value="litera"))
        each.body.append(WriteExcelNode(column="Status", value="manual"))
        return new_routine(SetVariableNode(name="Pret", value="1"), each)

    def test_columns_and_text_format(self):
        path = _tmp()
        self.paths.append(path)
        template.export_template(self._routine(), path, "4.22.0")
        wb = openpyxl.load_workbook(path)
        ws = wb.worksheets[0]
        self.assertEqual(["Id", "Pret", "Mesaj", "Status"], [c.value for c in ws[1]])
        self.assertEqual("@", ws.cell(row=2, column=1).number_format)
        self.assertEqual("@", ws.cell(row=500, column=1).number_format)
        self.assertIsNotNone(ws.cell(row=1, column=1).comment)
        self.assertEqual("A2", ws.freeze_panes)
        self.assertEqual(1, len(ws.data_validations.dataValidation))
        meta = wb["_autopilot"]
        self.assertEqual("hidden", meta.sheet_state)
        values = {meta.cell(row=r, column=1).value: meta.cell(row=r, column=2).value for r in range(1, meta.max_row + 1)}
        self.assertEqual("Test", values["routine"])
        self.assertEqual("4.22.0", values["app_version"])
        self.assertEqual("Id", values["key_column"])
        self.assertEqual(3, len(wb.sheetnames))

    def test_letter_key_header_row_and_errors(self):
        each = ForEachRowNode(key_column="A", status_column="", header_row=3, sheet="Comenzi")
        self.assertEqual("Nr", template.template_columns(ForEachRowNode(key_column="Nr"))[0])
        path = _tmp()
        self.paths.append(path)
        template.export_template(new_routine(each), path)
        ws = openpyxl.load_workbook(path)["Comenzi"]
        self.assertEqual("Id", ws.cell(row=3, column=1).value)
        self.assertIsNone(ws.cell(row=3, column=2).value)
        self.assertEqual("@", ws.cell(row=4, column=1).number_format)
        with self.assertRaises(RprStepException):
            template.export_template(new_routine(SetVariableNode(name="x")), _tmp())
        with self.assertRaises(RprStepException):
            template.export_template(new_routine(ForEachRowNode(header_row=0)), _tmp())

    def test_exported_template_runs_with_leading_zeros(self):
        path = _tmp()
        self.paths.append(path)
        template.export_template(self._routine(), path)
        wb = openpyxl.load_workbook(path)
        ws = wb.worksheets[0]
        for r, v in enumerate(["00123", "123456789012345678"], start=2):
            ws.cell(row=r, column=1, value=v)
        wb.save(path)
        engine, *_ = create()
        each = ForEachRowNode(file_path=path, key_column="Id", status_column="Status")
        each.body.append(WriteExcelNode(column="Pret", value="{Rand.Id}"))
        result = engine.run(new_routine(each))
        self.assertTrue(result.success, result.message)
        with ExcelTable.open(path, None, 1) as t:
            self.assertEqual("00123", t.get_cell(2, "Pret"))
            self.assertEqual("123456789012345678", t.get_cell(3, "Pret"))
            self.assertEqual("OK", t.get_cell(3, "Status"))


class KeyAliasesAndCheck(unittest.TestCase):
    def setUp(self):
        self.paths = []

    def tearDown(self):
        for p in self.paths:
            if os.path.exists(p):
                os.remove(p)

    def test_id_finds_ueid_only_for_key(self):
        path = _file([("A1", "x")], header=("Nume", " UEID ", "Cod material"))
        self.paths.append(path)
        with ExcelTable.open(path, None, 1) as t:
            self.assertEqual(2, t.resolve_column("Id", key=True))
            self.assertNotEqual(2, t.resolve_column("Id"))              # la scriere: nu (Id = litera ID)
            self.assertEqual(3, t.resolve_column("cod_material"))       # spatii / _ / litere mari
            self.assertEqual("UEID", t.title_of("ID"))
            t.write(2, "Cod", "SAP")                                   # nu peste UEID
            self.assertEqual("x", t.get_cell(2, "UEID"))

    def test_routine_with_id_runs_on_ueid_file(self):
        path = _file([("100", "a"), ("200", "b")])
        self.paths.append(path)
        engine, *_ = create()
        each = ForEachRowNode(file_path=path, key_column="Id", status_column="Status")
        each.body.append(WriteExcelNode(column="Rezultat", value="gasit {Rand.Id}"))
        result = engine.run(new_routine(each))
        self.assertTrue(result.success, result.message)
        self.assertEqual(2, result.rows_ok)
        with ExcelTable.open(path, None, 1) as t:
            self.assertEqual("gasit 200", t.get_cell(3, "Rezultat"))

    def test_missing_and_duplicate_ids_ask_before_first_row(self):
        rows = [("1", "a"), ("2", "b"), (None, "fara id"), ("1", "dublura")]
        for answer, runs in ((KEY_CHECK_BUTTONS[0], 2), (KEY_CHECK_BUTTONS[1], 0)):
            path = _file(rows)
            self.paths.append(path)
            engine, _, host, _ = create()
            host.answers = [answer]
            asked = []
            orig = host.ask
            host.ask = lambda title, text, buttons: (asked.append(text), orig(title, text, buttons))[1]
            each = ForEachRowNode(file_path=path, key_column="Id", status_column="Status")
            each.body.append(SetVariableNode(name="Nr", operation=VariableOperation.Increment))
            result = engine.run(new_routine(SetVariableNode(name="Nr", value="0"), each))
            self.assertEqual(1, len(asked))
            self.assertIn("4", asked[0])        # randul fara Id
            self.assertIn("\"1\"", asked[0])     # Id repetat
            self.assertIn("2, 5", asked[0])
            self.assertEqual(str(runs), result.final_variables.get("Nr", "0"), answer)

    def test_clean_file_does_not_ask(self):
        path = _file([("1", "a"), ("2", "b")])
        self.paths.append(path)
        engine, _, host, _ = create()
        host.answers = ["nu trebuia folosit"]
        each = ForEachRowNode(file_path=path, key_column="UEID", status_column="Status")
        result = engine.run(new_routine(each))
        self.assertTrue(result.success, result.message)
        self.assertEqual(["nu trebuia folosit"], host.answers)



class StudioButton(unittest.TestCase):
    """Studio real (Tk): 📤 Exporta sablon Excel din panoul buclei Excel."""
    @classmethod
    def setUpClass(cls):
        try:
            import tkinter as tk
            cls.root = tk.Tk()
            cls.root.geometry("1366x740")
        except Exception as ex:  # noqa: BLE001 — fara ecran (CI)
            raise unittest.SkipTest(f"fara afisaj: {ex}")
        import ui_theme
        from rpr.storage import RprPaths
        from test_views import COLORS
        cls.d = tempfile.mkdtemp(prefix="apdc_tpl_ui_")
        cls.RprPaths, cls.ui_theme = RprPaths, ui_theme
        cls.old_root, cls.old_aspect = RprPaths.root, ui_theme._SETTINGS
        RprPaths.root = os.path.join(cls.d, "Studio")
        ui_theme._SETTINGS = os.path.join(cls.d, "aspect.json")
        import studio_ui
        cls.studio_ui = studio_ui
        cls.p = studio_ui.StudioPanel(cls.root, COLORS)
        cls.p.pack(fill="both", expand=True)
        cls.root.update()

    @classmethod
    def tearDownClass(cls):
        cls.RprPaths.root, cls.ui_theme._SETTINGS = cls.old_root, cls.old_aspect
        cls.root.destroy()

    def _buttons(self, w, out):
        for ch in w.winfo_children():
            try:
                out.append((ch.cget("text"), ch))
            except Exception:  # noqa: BLE001
                pass
            self._buttons(ch, out)
        return out

    def test_export_button_saves_and_links_file(self):
        p, su = self.p, self.studio_ui
        each = ForEachRowNode(key_column="", status_column="Status")
        each.body.append(WriteExcelNode(column="Pret", value="1"))
        p.insert_node(each)
        p.editor.show(each)
        self.root.update()
        btn = dict(self._buttons(p.props.inner, [])).get("📤 Exporta sablon Excel")
        self.assertIsNotNone(btn)
        out = os.path.join(self.d, "Sablon.xlsx")
        opened = []
        orig = (su.filedialog.asksaveasfilename, su.messagebox.askyesno, p.open_path)
        su.filedialog.asksaveasfilename = lambda **k: out
        su.messagebox.askyesno = lambda *a, **k: True
        p.open_path = opened.append
        try:
            btn.invoke()
            self.root.update()
        finally:
            su.filedialog.asksaveasfilename, su.messagebox.askyesno, p.open_path = orig
        self.assertEqual([out], opened)
        ws = openpyxl.load_workbook(out).worksheets[0]
        self.assertEqual(["Id", "Pret", "Status"], [c.value for c in ws[1]])
        self.assertEqual(os.path.normpath(out), each.file_path)
        self.assertTrue(p.dirty)


if __name__ == "__main__":
    unittest.main()
