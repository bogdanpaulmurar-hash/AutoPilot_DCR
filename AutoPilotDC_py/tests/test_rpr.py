# -*- coding: utf-8 -*-
# ============================================================
# AutoPilot DC v4.1 — Teste pentru motorul RPR (Python)
# Aceleasi scenarii ca testele C# (RprEngineTests + SafetyTests).
# Un driver fals simuleaza ecranul — nu e nevoie de browser / SAP.
#
# Rulare:  python -m unittest discover -s tests      (din folderul aplicatiei)
# ============================================================

import json
import os
import re
import sys
import tempfile
import threading
import unittest
import uuid

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from rpr import (ActionNode, AlertDecision, BreakNode, CellColor, CompareOperator, Condition, ConditionKind,  # noqa: E402
                 DriverProvider, ErrorDecision, ErrorPolicy, ForEachRowNode, IfNode, LegacyConverter,
                 LegacyStepNode, Locator, LocatorKind, LogNode, LoopNode, ObjectRepository, QuestionNode,
                 Routine, RprAction, RprEngine, RprStepException, RunAlert, RunGuard, RunHost, SetVariableNode,
                 TryNode, UiDriver, UiObject, UiTechnology, VariableOperation, VariableStore, WriteExcelNode,
                 matches_expected)
from rpr.conditions import compare, try_number  # noqa: E402
from rpr.drivers.base import NOT_FOUND, ElementNotFound, ElementState  # noqa: E402
from rpr.excel import ExcelTable  # noqa: E402
from rpr.model import node_from_dict  # noqa: E402
from rpr.variables import format_net_date  # noqa: E402


# ── Driver fals: "ecranul" e un dictionar nume obiect -> text ──
class FakeDriver(UiDriver):
    technology = UiTechnology.Web
    display_name = "Fake"

    def __init__(self):
        self.screen = {}
        self.actions = []
        self.failures_before_success = 0

    def _has(self, name):
        return name.lower() in {k.lower() for k in self.screen}

    def _get(self, name):
        return next(v for k, v in self.screen.items() if k.lower() == name.lower())

    def get_state(self, obj, timeout_ms, stop):
        if self._has(obj.name):
            t = self._get(obj.name)
            return ElementState(True, True, True, t, t, obj.locators[0] if obj.locators else None)
        return NOT_FOUND

    def execute(self, r, stop):
        if r.target is not None and not self._has(r.target.name):
            raise ElementNotFound(r.target, r.timeout_ms)
        if self.failures_before_success > 0:
            self.failures_before_success -= 1
            raise RprStepException("esec simulat")
        self.actions.append(f"{r.action}:{r.target.name if r.target else ''}:{r.value}")
        if r.action == RprAction.TypeText and r.target is not None:
            self.screen[r.target.name] = r.value
        if r.action in (RprAction.ReadText, RprAction.ReadValue) and r.target is not None:
            return self._get(r.target.name)
        return None


class FakeHost(RunHost):
    def __init__(self):
        self.entries = []
        self.answers = []
        self.on_error = ErrorDecision.STOP
        self.alert_decision = AlertDecision.CONTINUE
        self.alerts = []

    def log(self, entry):
        self.entries.append(entry)

    def ask(self, title, text, buttons):
        return self.answers.pop(0) if self.answers else buttons[0]

    def ask_on_error(self, node, error, inside_excel_loop):
        return self.on_error

    def ask_alert(self, node, alert, inside_excel_loop):
        self.alerts.append(alert)
        return self.alert_decision


class FakeGuard(RunGuard):
    def __init__(self, alerts_left):
        self.alerts_left = alerts_left

    def check_before_action(self, node):
        self.alerts_left -= 1
        return RunAlert("A aparut o fereastra de eroare", "Eroare: test", "window:Eroare") if self.alerts_left >= 0 else None


def create():
    repo = ObjectRepository()
    for name, loc in (("Camp", "camp"), ("Buton", "btn"), ("Lipsa", "lipsa")):
        repo.add(UiObject(name=name, locators=[Locator(LocatorKind.Id, loc)]))
    driver = FakeDriver()
    driver.screen["Camp"] = ""
    driver.screen["Buton"] = "OK"
    host = FakeHost()
    provider = DriverProvider({t: (lambda d=driver: d) for t in UiTechnology})
    engine = RprEngine(repo, provider, host, os.path.join(tempfile.gettempdir(), "apdc-teste-py"))
    return engine, driver, host, repo


def new_routine(*steps):
    r = Routine(name="Test")
    r.settings.highlight_before_action = False
    r.settings.auto_retries = 0
    r.settings.screenshot_on_error = False
    r.settings.default_error_policy = ErrorPolicy.Stop
    r.steps.extend(steps)
    return r


class VariablesAndCompare(unittest.TestCase):
    def test_resolve_known_unknown_escapes(self):
        v = VariableStore()
        v.set("Nume", "Ana Maria")
        v.set("Rand.IBase", "12345")
        self.assertEqual("Salut Ana Maria / 12345", v.resolve("Salut {Nume} / {Rand.IBase}"))
        unknown = []
        self.assertEqual("{Necunoscut}", v.resolve("{Necunoscut}", unknown))
        self.assertEqual(1, len(unknown))
        self.assertEqual("{literal}", v.resolve("{{literal}}"))
        self.assertEqual("Ana%20Maria", v.resolve("{Nume:url}"))
        self.assertEqual("ANA MARIA", v.resolve("{nume:upper}"))
        self.assertRegex(v.resolve("{Azi}"), r"^\d{2}\.\d{2}\.\d{4}$")
        self.assertRegex(v.resolve("{Acum:yyyyMMdd_HHmm}"), r"^\d{8}_\d{4}$")

    def test_net_date_format(self):
        import datetime as dt
        d = dt.datetime(2026, 3, 7, 14, 5, 9, 123000)
        self.assertEqual("07.03.2026 14:05:09", format_net_date(d, "dd.MM.yyyy HH:mm:ss"))
        self.assertEqual("20260307_1405.123", format_net_date(d, "yyyyMMdd_HHmm.fff"))

    def test_numbers(self):
        for text, expected in (("1.234,5", 1234.5), ("1,234.5", 1234.5), ("12,5", 12.5), ("1,234", 1234), ("-7", -7)):
            ok, d = try_number(text)
            self.assertTrue(ok, text)
            self.assertEqual(float(d), expected, text)

    def test_compare_numeric_when_possible(self):
        self.assertTrue(compare("10", CompareOperator.GreaterThan, "9", True))
        self.assertTrue(compare("abc", CompareOperator.Contains, "B", True))
        self.assertFalse(compare("abc", CompareOperator.Contains, "B", False))
        self.assertTrue(compare("5,0", CompareOperator.Equals, "5", True))
        self.assertTrue(compare("  ", CompareOperator.IsEmpty, "", True))


class Flow(unittest.TestCase):
    def test_loop_if_break_variables(self):
        engine, *_ = create()
        loop = LoopNode(count="10", counter_variable="i")
        loop.body.append(SetVariableNode(name="Suma", operation=VariableOperation.Increment, value="{i}"))
        stop = IfNode(condition=Condition(kind=ConditionKind.Compare, left="{i}", operator=CompareOperator.Equals, right="4"))
        stop.then.append(BreakNode())
        loop.body.append(stop)
        result = engine.run(new_routine(SetVariableNode(name="Suma", value="0"), loop))
        self.assertTrue(result.success, result.message)
        self.assertEqual("10", result.final_variables["Suma"])  # 1+2+3+4

    def test_action_type_and_read_with_variables(self):
        engine, driver, *_ = create()
        r = new_routine(ActionNode(action=RprAction.TypeText, target="Camp", value="Cod-{X}"),
                        ActionNode(action=RprAction.ReadText, target="Camp", output_variable="Citit"))
        r.variables["X"] = "42"
        result = engine.run(r)
        self.assertTrue(result.success, result.message)
        self.assertEqual("Cod-42", result.final_variables["Citit"])
        self.assertIn("TypeText:Camp:Cod-42", driver.actions)

    def test_missing_object_fails_before_running(self):
        engine, driver, *_ = create()
        result = engine.run(new_routine(ActionNode(action=RprAction.Click, target="NuExista")))
        self.assertFalse(result.success)
        self.assertIn("NuExista", result.message)
        self.assertEqual([], driver.actions)

    def test_try_catches_error_and_continues(self):
        engine, *_ = create()
        t = TryNode(error_variable="Err")
        t.body.append(ActionNode(action=RprAction.Click, target="Lipsa", timeout_ms=10))
        t.catch.append(SetVariableNode(name="Prins", value="da"))
        result = engine.run(new_routine(t, SetVariableNode(name="Dupa", value="ok")))
        self.assertTrue(result.success, result.message)
        self.assertEqual("da", result.final_variables["Prins"])
        self.assertEqual("ok", result.final_variables["Dupa"])
        self.assertIn("Lipsa", result.final_variables["Err"])

    def test_error_policy_continue_and_stop(self):
        engine, *_ = create()
        cont = engine.run(new_routine(
            ActionNode(action=RprAction.Click, target="Lipsa", timeout_ms=10, on_error=ErrorPolicy.Continue),
            SetVariableNode(name="Ajuns", value="da")))
        self.assertTrue(cont.success)
        self.assertEqual(1, cont.errors)
        self.assertEqual("da", cont.final_variables["Ajuns"])

        engine2, *_ = create()
        stop = engine2.run(new_routine(ActionNode(action=RprAction.Click, target="Lipsa", timeout_ms=10),
                                       SetVariableNode(name="Ajuns", value="da")))
        self.assertFalse(stop.success)
        self.assertNotIn("Ajuns", stop.final_variables)

    def test_ask_on_error_retry_then_skip(self):
        engine, driver, host, _ = create()
        driver.failures_before_success = 1
        r = new_routine(ActionNode(action=RprAction.Click, target="Buton", on_error=ErrorPolicy.Ask))
        host.on_error = ErrorDecision.RETRY
        result = engine.run(r)
        self.assertTrue(result.success, result.message)
        self.assertEqual(0, result.errors)

    def test_auto_retry_recovers(self):
        engine, driver, *_ = create()
        driver.failures_before_success = 1
        r = new_routine(ActionNode(action=RprAction.Click, target="Buton"))
        r.settings.auto_retries = 2
        result = engine.run(r)
        self.assertTrue(result.success, result.message)
        self.assertEqual(0, result.errors)

    def test_question_stop_button(self):
        engine, _, host, _ = create()
        host.answers.append("Opreste")
        result = engine.run(new_routine(QuestionNode(text="?", buttons="Continua,Opreste"),
                                        SetVariableNode(name="Ajuns", value="da")))
        self.assertTrue(result.success)  # oprire voita, nu eroare
        self.assertNotIn("Ajuns", result.final_variables)

    def test_condition_element_exists_and_text(self):
        engine, *_ = create()
        i = IfNode(condition=Condition(kind=ConditionKind.AllOf, children=[
            Condition(kind=ConditionKind.ElementExists, target="Buton"),
            Condition(kind=ConditionKind.ElementText, target="Buton", operator=CompareOperator.Equals, right="ok"),
            Condition(kind=ConditionKind.ElementExists, target="Lipsa", negate=True),
        ]))
        i.then.append(SetVariableNode(name="Ramura", value="Atunci"))
        i.else_.append(SetVariableNode(name="Ramura", value="Altfel"))
        result = engine.run(new_routine(i))
        self.assertEqual("Atunci", result.final_variables["Ramura"])

    def test_stop_from_other_thread(self):
        engine, *_ = create()
        loop = LoopNode(count="1000")
        loop.body.append(LogNode(message="x"))
        from rpr.model import WaitNode
        loop.body.append(WaitNode(seconds=0.05))
        threading.Timer(0.3, engine.stop).start()
        result = engine.run(new_routine(loop))
        self.assertTrue(result.cancelled)
        self.assertFalse(result.success)


class Excel(unittest.TestCase):
    def _make_file(self):
        import openpyxl
        path = os.path.join(tempfile.gettempdir(), f"apdc_test_{uuid.uuid4().hex}.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Date"
        for r, v in enumerate(["Cod", "A1", "RAU", "C3"], 1):
            ws.cell(row=r, column=1, value=v)
        ws.cell(row=1, column=5, value="Formula")
        ws.cell(row=2, column=5, value="=1+1")
        wb.save(path)
        return path

    def test_for_each_row_writes_results_status_and_skips_failed_row(self):
        path = self._make_file()
        try:
            engine, *_ = create()
            each = ForEachRowNode(file_path=path, key_column="Cod", status_column="Status")
            bad = IfNode(condition=Condition(kind=ConditionKind.Compare, left="{Rand.Cod}", operator=CompareOperator.Equals, right="RAU"))
            bad.then.append(ActionNode(action=RprAction.Click, target="Lipsa", timeout_ms=10, on_error=ErrorPolicy.SkipRow))
            each.body.append(bad)
            each.body.append(WriteExcelNode(column="Rezultat", value="procesat {Rand.Cod} (rand {Rand.#})", color=CellColor.Blue))

            result = engine.run(new_routine(each))
            self.assertTrue(result.success, result.message)
            self.assertEqual(2, result.rows_ok)
            self.assertEqual(1, result.rows_failed)

            with ExcelTable.open(path, None, 1) as check:
                self.assertEqual("procesat A1 (rand 2)", check.get_cell(2, "Rezultat"))
                self.assertEqual("OK", check.get_cell(2, "Status"))
                self.assertTrue(check.get_cell(3, "Status").startswith("EROARE"))
                self.assertEqual("", check.get_cell(3, "Rezultat"))
                self.assertEqual("procesat C3 (rand 4)", check.get_cell(4, "Rezultat"))
            import openpyxl
            ws = openpyxl.load_workbook(path)["Date"]
            self.assertEqual("=1+1", ws.cell(row=2, column=5).value)       # formula pastrata
            self.assertEqual("FFC7CE", ws.cell(row=3, column=1).fill.start_color.rgb[-6:])  # rand rosu
        finally:
            os.remove(path)

    def test_resume_skips_completed_rows(self):
        path = self._make_file()
        try:
            engine, *_ = create()
            each = ForEachRowNode(file_path=path, key_column="Cod")
            each.body.append(SetVariableNode(name="Nr", operation=VariableOperation.Increment))
            r1 = engine.run(new_routine(SetVariableNode(name="Nr", value="0"), each))
            self.assertEqual(3, r1.rows_ok)
            engine2, *_ = create()
            r2 = engine2.run(new_routine(SetVariableNode(name="Nr", value="0"), each))
            self.assertEqual(3, r2.rows_skipped)
            self.assertEqual("0", r2.final_variables["Nr"])
        finally:
            os.remove(path)

    def test_unfinished_row_asks_redo_mark_ok_or_stop(self):
        from rpr.engine import IN_PROGRESS, UNFINISHED_BUTTONS
        for answer, expected_runs, status_ok in ((UNFINISHED_BUTTONS[0], 3, "OK"),
                                                 (UNFINISHED_BUTTONS[1], 2, "OK (verificat manual)"),
                                                 (UNFINISHED_BUTTONS[2], 1, IN_PROGRESS)):
            path = self._make_file()
            try:
                # prima rulare: randul 2 OK, randul 3 ramas "IN LUCRU" (aplicatia s-a inchis)
                with ExcelTable.open(path, None, 1) as t:
                    t.write(2, "Status", "OK")
                    t.write(3, "Status", f"{IN_PROGRESS} 29.10 10:00")
                    t.save()
                engine, _, host, _ = create()
                host.answers = [answer]
                each = ForEachRowNode(file_path=path, key_column="Cod", status_column="Status")
                each.body.append(SetVariableNode(name="Nr", operation=VariableOperation.Increment))
                r = engine.run(new_routine(SetVariableNode(name="Nr", value="1"), each))
                self.assertEqual(str(expected_runs), r.final_variables["Nr"], answer)
                with ExcelTable.open(path, None, 1) as check:
                    self.assertTrue(check.get_cell(3, "Status").startswith(status_ok), check.get_cell(3, "Status"))
                    if answer != UNFINISHED_BUTTONS[2]:
                        self.assertEqual("OK", check.get_cell(4, "Status"))
                    else:
                        self.assertEqual("", check.get_cell(4, "Status"))
            finally:
                os.remove(path)

    def test_locked_or_missing_file_gives_clear_message(self):
        engine, *_ = create()
        result = engine.run(new_routine(ForEachRowNode(file_path=r"C:\nu\exista.xlsx")))
        self.assertFalse(result.success)
        self.assertIn("nu exista", result.message)


class Safety(unittest.TestCase):
    def _read(self, expected):
        return ActionNode(action=RprAction.ReadText, target="Camp", output_variable="Mesaj", expected_value=expected)

    def _engine(self):
        engine, driver, host, repo = create()
        driver.screen["Camp"] = "Comanda salvata 4711"
        return engine, driver, host

    def test_matching(self):
        cases = [("Comanda salvata 4711", "comanda salvata 4711", True), ("Comanda salvata 4711", "Comanda salvata*", True),
                 ("Comanda salvata 4711", "*4711", True), ("S: Document 12 salvat", "S:*", True),
                 ("E: Material inexistent", "S:*", False), ("abc", "abd", False), ("a.b", "a*b", True), ("axb", "a.b", False)]
        for actual, expected, match in cases:
            self.assertEqual(match, matches_expected(actual, expected), (actual, expected))

    def test_expected_value_matches_no_question(self):
        engine, _, host = self._engine()
        result = engine.run(new_routine(self._read("Comanda salvata*")))
        self.assertTrue(result.success, result.message)
        self.assertEqual([], host.alerts)

    def test_expected_value_differs_asks_and_stops(self):
        engine, _, host = self._engine()
        host.alert_decision = AlertDecision.STOP
        result = engine.run(new_routine(self._read("Eroare*"), LogNode(message="nu trebuie")))
        self.assertFalse(result.success)
        self.assertEqual(1, len(host.alerts))
        self.assertIn("difera", host.alerts[0].title)
        self.assertEqual("Comanda salvata 4711", result.final_variables["Mesaj"])
        self.assertEqual(0, result.steps_executed)

    def test_expected_value_differs_continue(self):
        engine, _, host = self._engine()
        result = engine.run(new_routine(self._read("Altceva"), SetVariableNode(name="Dupa", value="da")))
        self.assertTrue(result.success, result.message)
        self.assertEqual("da", result.final_variables["Dupa"])

    def test_expected_value_inside_try_goes_to_catch(self):
        engine, _, host = self._engine()
        t = TryNode(error_variable="Eroare")
        t.body.append(self._read("Altceva"))
        t.catch.append(SetVariableNode(name="Prins", value="da"))
        result = engine.run(new_routine(t))
        self.assertTrue(result.success, result.message)
        self.assertEqual("da", result.final_variables["Prins"])
        self.assertEqual([], host.alerts)

    def test_guard_ignore_similar_asks_once(self):
        engine, driver, host = self._engine()
        engine.guard = FakeGuard(5)
        host.alert_decision = AlertDecision.IGNORE_SIMILAR
        result = engine.run(new_routine(ActionNode(action=RprAction.Click, target="Camp"),
                                        ActionNode(action=RprAction.Click, target="Camp")))
        self.assertTrue(result.success, result.message)
        self.assertEqual(1, len(host.alerts))
        self.assertEqual(2, len(driver.actions))

    def test_guard_disabled_in_routine(self):
        engine, _, host = self._engine()
        engine.guard = FakeGuard(5)
        r = new_routine(ActionNode(action=RprAction.Click, target="Camp"))
        r.settings.guard_enabled = False
        self.assertTrue(engine.run(r).success)
        self.assertEqual([], host.alerts)

    def test_guard_stop_before_action(self):
        engine, driver, host = self._engine()
        engine.guard = FakeGuard(1)
        host.alert_decision = AlertDecision.STOP
        result = engine.run(new_routine(ActionNode(action=RprAction.Click, target="Camp")))
        self.assertFalse(result.success)
        self.assertEqual([], driver.actions)


class Persistence(unittest.TestCase):
    def test_json_round_trip_keeps_tree_and_types(self):
        i = IfNode(condition=Condition(kind=ConditionKind.ElementVisible, target="Buton"))
        i.then.append(ActionNode(action=RprAction.Click, target="Buton"))
        i.else_.append(LoopNode(count="2"))
        r = new_routine(i, ForEachRowNode(file_path="x.xlsx"))
        r.variables["A"] = "1"
        text = json.dumps(r.to_dict(), ensure_ascii=False, indent=2)
        self.assertIn('"kind": "if"', text)
        self.assertIn('"else": [', text)
        self.assertTrue(text.index('"kind"') < text.index('"id"'))  # "kind" primul (cerut de C#)
        back = Routine.from_dict(json.loads(text))
        bi = back.steps[0]
        self.assertIsInstance(bi, IfNode)
        self.assertEqual(ConditionKind.ElementVisible, bi.condition.kind)
        self.assertIsInstance(bi.then[0], ActionNode)
        self.assertIsInstance(bi.else_[0], LoopNode)
        self.assertIsInstance(back.steps[1], ForEachRowNode)
        self.assertEqual("1", back.variables["A"])

    def test_unknown_fields_are_kept(self):
        d = {"kind": "action", "id": "x1", "action": "Click", "target": "B", "camp_viitor": 7}
        n = node_from_dict(d)
        self.assertEqual(7, n.to_dict()["camp_viitor"])

    def test_legacy_converter(self):
        repo = ObjectRepository()
        r = LegacyConverter.convert("Vechi", [
            {"step_type": "browser_click", "key_combo": "#submit", "description": "Buton Submit"},
            {"step_type": "CLICK", "x": 100, "y": 200, "delay_before": 1.5},
            {"step_type": "QUESTION", "question_text": "Continui?"},
        ], repo)
        self.assertEqual(1, len(repo.objects))
        self.assertIsInstance(r.steps[0], ActionNode)
        self.assertEqual(repo.objects[0].name, r.steps[0].target)
        self.assertIsInstance(r.steps[1], LegacyStepNode)
        self.assertEqual(1500, r.steps[1].delay_before_ms)
        self.assertEqual(0, r.steps[1].step["delay_before"])
        self.assertIsInstance(r.steps[2], QuestionNode)

    def test_repository_unique_names(self):
        repo = ObjectRepository()
        a = repo.add(UiObject(name="Buton OK"))
        b = repo.add(UiObject(name="Buton OK"))
        self.assertEqual("Buton_OK", a.name)
        self.assertEqual("Buton_OK_2", b.name)
        self.assertFalse(repo.rename(b, "Buton_OK"))
        self.assertTrue(repo.rename(b, "Buton Anuleaza"))
        self.assertEqual("Buton_Anuleaza", b.name)


if __name__ == "__main__":
    unittest.main()
