# -*- coding: utf-8 -*-
# ============================================================
# AutoPilot DC — Teste: managerul de rutine si pregatirile (rpr/batch.py)
#   Rulare:  py -m unittest tests.test_batch
# ============================================================

import os
import sys
import tempfile
import unittest

APP = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, APP)

import openpyxl  # noqa: E402

from rpr import DriverProvider, ObjectRepository, RprEngine, RunHost  # noqa: E402
from rpr.batch import (PREP_EXT, Preparation, PreparationRunner, PrepItem, RoutineLibrary)  # noqa: E402
from rpr.model import LogNode, LoopNode, Routine, SetVariableNode, StopNode  # noqa: E402
from rpr.storage import RoutineStore  # noqa: E402


def _save(folder, name, steps, description=""):
    r = Routine(name=name, description=description)
    r.settings.do_not_disturb = False
    r.settings.guard_enabled = False
    r.steps = steps
    path = os.path.join(folder, name + ".rpr.json")
    RoutineStore.save(r, path)
    return path


def _engine_run_one(out):
    def run_one(routine, item, folder):
        return RprEngine(ObjectRepository(), DriverProvider(), RunHost(), out).run(routine), ""
    return run_one


class LibraryTests(unittest.TestCase):
    def test_scan_lists_routines_preparations_and_broken_files(self):
        d = tempfile.mkdtemp(prefix="apdc_lib_")
        os.makedirs(os.path.join(d, "Sandbox"))
        _save(d, "A_useri", [LogNode(message="x"), LoopNode(count="2", body=[LogNode(message="y")])], "Creeaza useri")
        _save(os.path.join(d, "Sandbox"), "B_roluri", [LogNode(message="z")])
        with open(os.path.join(d, "stricata.rpr.json"), "w", encoding="utf-8") as f:
            f.write("{nu e json")
        Preparation(name="Sandbox", items=[PrepItem("A_useri.rpr.json")]).save(os.path.join(d, "Sandbox" + PREP_EXT))
        lib = RoutineLibrary().scan(d)
        names = {r.name: r for r in lib.routines}
        self.assertEqual({"A_useri", "B_roluri", "stricata"}, set(names))
        self.assertEqual(3, names["A_useri"].steps)
        self.assertEqual("Creeaza useri", names["A_useri"].description)
        self.assertTrue(names["stricata"].error)
        self.assertEqual(1, len(lib.preparations))
        self.assertIsNotNone(lib.find(os.path.join(d, "Sandbox", "B_roluri.rpr.json")))

    def test_missing_folder_is_empty(self):
        self.assertEqual([], RoutineLibrary().scan(os.path.join(tempfile.gettempdir(), "nu_exista_xyz")).routines)


class PreparationTests(unittest.TestCase):
    def test_save_load_relative_paths(self):
        d = tempfile.mkdtemp(prefix="apdc_prep_")
        a = _save(d, "A", [LogNode(message="x")])
        p = Preparation(name="Sandbox", path=os.path.join(d, "Sandbox" + PREP_EXT))
        p.items.append(PrepItem(p.relative(a), stop_on_error=False, variables={"Nr": "5"}))
        p.save()
        q = Preparation.load(p.path)
        self.assertEqual("A.rpr.json", q.items[0].path)        # relativa: folderul poate fi mutat
        self.assertEqual(os.path.normpath(a), q.resolve(q.items[0]))
        self.assertFalse(q.items[0].stop_on_error)
        self.assertEqual({"Nr": "5"}, q.items[0].variables)

    def test_run_in_order_with_reports_summary_and_documentation(self):
        d = tempfile.mkdtemp(prefix="apdc_prep_")
        out = os.path.join(d, "rezultate")
        a = _save(d, "01_Useri", [LoopNode(count="{Nr}", counter_variable="i",
                                           body=[SetVariableNode(name="User", value="user{i}")])], "Creeaza useri")
        b = _save(d, "02_Roluri", [StopNode(message="rol lipsa", as_error=True)])
        c = _save(d, "03_Obiecte", [LogNode(message="obiecte")])
        p = Preparation(name="Sandbox test", output_folder=out, path=os.path.join(d, "S" + PREP_EXT))
        p.items = [PrepItem(p.relative(a), variables={"Nr": "3"}),
                   PrepItem(p.relative(b), stop_on_error=False),
                   PrepItem(p.relative(c)),
                   PrepItem("lipsa.rpr.json", enabled=False)]
        seen = []
        res = PreparationRunner(p, _engine_run_one(d), d, on_progress=lambda i, n, name: seen.append((i, n))).run()
        self.assertEqual([(1, 3), (2, 3), (3, 3)], seen)
        self.assertEqual(["OK", "EROARE", "OK"], [s.status for s in res.steps])
        self.assertFalse(res.success)
        self.assertTrue(res.folder.startswith(out))
        # fiecare rutina are raportul ei in folderul pregatirii (raportul pornit de pregatire)
        self.assertTrue(all(s.report.startswith(os.path.join(res.folder, "rapoarte")) for s in res.steps), res.steps)
        rows = list(openpyxl.load_workbook(res.steps[0].report)["Rezultate"].iter_rows(values_only=True))
        self.assertEqual(4, len(rows))                          # 3 repetari = 3 useri
        # rezumatul + definitiile
        wb = openpyxl.load_workbook(res.summary)
        self.assertEqual(["Rezumat", "Definitii"], wb.sheetnames)
        values = [c for row in wb["Rezumat"].iter_rows(values_only=True) for c in row if c]
        self.assertIn("02_Roluri", values)
        self.assertIn("rol lipsa", values)
        with open(res.documentation, encoding="utf-8") as f:
            doc = f.read()
        self.assertIn("Creeaza useri", doc)
        self.assertIn("Repeta de {Nr} ori", doc)
        # rutina din disc nu a fost modificata (raportul e pornit doar in copie)
        self.assertNotIn("report", RoutineStore.load(a).settings.extra)

    def test_stop_on_error_skips_the_rest(self):
        d = tempfile.mkdtemp(prefix="apdc_prep_")
        a = _save(d, "A", [StopNode(message="gresit", as_error=True)])
        b = _save(d, "B", [LogNode(message="x")])
        p = Preparation(name="P", report=False, video=False, path=os.path.join(d, "P" + PREP_EXT))
        p.items = [PrepItem(p.relative(a)), PrepItem(p.relative(b))]
        res = PreparationRunner(p, _engine_run_one(d), d).run()
        self.assertEqual(["EROARE", "SARIT"], [s.status for s in res.steps])
        self.assertEqual("", res.steps[0].report)

    def test_missing_routine_and_cancel(self):
        d = tempfile.mkdtemp(prefix="apdc_prep_")
        b = _save(d, "B", [LogNode(message="x")])
        p = Preparation(name="P", report=False, path=os.path.join(d, "P" + PREP_EXT))
        p.items = [PrepItem("nu_exista.rpr.json", stop_on_error=False), PrepItem(p.relative(b)),
                   PrepItem(p.relative(b))]
        runner = PreparationRunner(p, None, d)

        def run_one(routine, item, folder):
            runner.cancel()          # STOP in timpul primei rutine rulate
            return RprEngine(ObjectRepository(), DriverProvider(), RunHost(), d).run(routine), ""
        runner._run_one = run_one
        res = runner.run()
        self.assertEqual(["LIPSA", "OK", "SARIT"], [s.status for s in res.steps])
        self.assertTrue(res.cancelled)
        self.assertTrue(os.path.isfile(res.documentation))


class CliTests(unittest.TestCase):
    def test_run_preparation_headless(self):
        from rpr import __main__ as cli
        d = tempfile.mkdtemp(prefix="apdc_cli_")
        a = _save(d, "A", [LogNode(message="x")])
        p = Preparation(name="CLI", video=False, output_folder=os.path.join(d, "out"),
                        path=os.path.join(d, "CLI" + PREP_EXT))
        p.items = [PrepItem(p.relative(a))]
        p.save()
        code, name, msg, summary = cli._run_preparation(p.path, False, __import__("datetime").datetime.now())
        self.assertEqual(0, code, msg)
        self.assertEqual("CLI", name)
        self.assertTrue(os.path.isfile(summary))


if __name__ == "__main__":
    unittest.main()
