# -*- coding: utf-8 -*-
# ============================================================
# AutoPilot DC — Teste: raportul Excel al rularii (rpr/report.py)
#   Rulare:  py -m unittest tests.test_report
# ============================================================

import os
import sys
import tempfile
import unittest

APP = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, APP)

import openpyxl  # noqa: E402

from rpr import DriverProvider, ObjectRepository, RprEngine, RunHost  # noqa: E402
from rpr.model import (ForEachRowNode, IfNode, LogNode, LoopNode, Routine, SetVariableNode, StopNode,  # noqa: E402
                       WriteExcelNode)
from rpr.model import Condition, ConditionKind, CompareOperator  # noqa: E402
from rpr.report import ReportSettings, has_variable, toggle_variable  # noqa: E402


def _routine(name="Raport test"):
    r = Routine(name=name)
    r.settings.do_not_disturb = False
    r.settings.guard_enabled = False
    return r


class _VideoHost(RunHost):
    class _Rec:
        active = True
        path = os.path.join(tempfile.gettempdir(), "film_test.apvid")

        def elapsed(self):
            return 75.0
    video_recorder = _Rec()


def _rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


class ReportSettingsTests(unittest.TestCase):
    def test_roundtrip_in_extra_and_json(self):
        from rpr.model import Routine as R
        r = _routine()
        self.assertFalse(ReportSettings.of(r).enabled)
        toggle_variable(r, "{Pret}", True)
        self.assertTrue(ReportSettings.of(r).enabled)          # bifa pe pas porneste raportul
        self.assertTrue(has_variable(r, "Pret"))
        r2 = R.from_dict(r.to_dict())                           # se pastreaza in .rpr.json (settings.report)
        self.assertEqual(["Pret"], ReportSettings.of(r2).variables)
        self.assertIn("report", r.to_dict()["settings"])
        toggle_variable(r2, "pret", False)
        self.assertFalse(has_variable(r2, "Pret"))

    def test_disabled_writes_nothing(self):
        d = tempfile.mkdtemp(prefix="apdc_rep_")
        r = _routine()
        r.steps.append(LogNode(message="x"))
        res = RprEngine(ObjectRepository(), DriverProvider(), RunHost(), d).run(r)
        self.assertEqual("", res.report_path)
        self.assertFalse(os.path.isdir(os.path.join(d, "rapoarte")))


class ReportRunTests(unittest.TestCase):
    def test_no_loop_one_row_with_values_steps_and_definition(self):
        d = tempfile.mkdtemp(prefix="apdc_rep_")
        r = _routine()
        r.steps += [SetVariableNode(name="Client", value="ACME"),
                    IfNode(condition=Condition(kind=ConditionKind.Compare, left="{Client}",
                                               operator=CompareOperator.Equals, right="ACME"),
                           then=[SetVariableNode(name="Rezultat", value="gasit")]),
                    LogNode(message="gata")]
        ReportSettings(enabled=True, variables=["Client", "{Client}-{Rezultat}"]).save_to(r)
        res = RprEngine(ObjectRepository(), DriverProvider(), _VideoHost(), d).run(r)
        self.assertTrue(res.success, res.message)
        self.assertTrue(os.path.isfile(res.report_path), res.report_path)
        wb = openpyxl.load_workbook(res.report_path)
        self.assertEqual(["Rezultate", "Pasi", "Rutina"], wb.sheetnames)
        rows = _rows(wb["Rezultate"])
        head = rows[0]
        self.assertEqual(2, len(rows))
        self.assertEqual("OK", rows[1][head.index("Status")])
        self.assertEqual("ACME", rows[1][head.index("Client")])
        self.assertEqual("ACME-gasit", rows[1][head.index("{Client}-{Rezultat}")])
        self.assertIn("gata", rows[1][head.index("Pasi parcursi")])
        self.assertEqual("film_test.apvid", rows[1][head.index("Film")])
        # foaia Pasi: si regula (Daca) apare, in ordinea executiei
        steps = [x[2].strip() for x in _rows(wb["Pasi"])[1:]]
        self.assertTrue(steps[0].startswith("{Client}") or "Client" in steps[0], steps)
        self.assertTrue(any(s.startswith("Daca") for s in steps), steps)
        definition = [x[1] for x in _rows(wb["Rutina"]) if x[1]]
        self.assertTrue(any("[Altfel]" in str(x) for x in definition), definition)

    def test_loop_iterations_are_items_and_stop_is_reported(self):
        d = tempfile.mkdtemp(prefix="apdc_rep_")
        r = _routine()
        r.steps.append(LoopNode(count="3", counter_variable="i", body=[
            LoopNode(count="2", body=[LogNode(message="interior")]),     # bucla din interior: nu e element nou
            IfNode(condition=Condition(kind=ConditionKind.Compare, left="{i}", operator=CompareOperator.Equals,
                                       right="3"),
                   then=[StopNode(message="oprit la 3", as_error=True)])]))
        ReportSettings(enabled=True, variables=["i"], steps_sheet=False).save_to(r)
        res = RprEngine(ObjectRepository(), DriverProvider(), RunHost(), d).run(r)
        self.assertFalse(res.success)
        wb = openpyxl.load_workbook(res.report_path)
        self.assertEqual(["Rezultate", "Rutina"], wb.sheetnames)
        rows = _rows(wb["Rezultate"])
        head = rows[0]
        self.assertEqual(["Repetarea 1", "Repetarea 2", "Repetarea 3"], [x[1] for x in rows[1:]])
        self.assertEqual(["OK", "OK", "EROARE"], [x[head.index("Status")] for x in rows[1:]])
        self.assertEqual("oprit la 3", rows[3][head.index("Mesaj")])
        self.assertEqual(["1", "2", "3"], [x[head.index("i")] for x in rows[1:]])

    def test_excel_rows_with_key_and_custom_folder(self):
        d = tempfile.mkdtemp(prefix="apdc_rep_")
        xl = os.path.join(d, "date.xlsx")
        wb = openpyxl.Workbook()
        wb.active.append(["ID", "Valoare"])
        for row in (["A1", "5"], ["A2", ""], ["A3", "7"]):
            wb.active.append(row)
        wb.save(xl)
        r = _routine()
        r.steps.append(ForEachRowNode(file_path=xl, key_column="ID", body=[
            IfNode(condition=Condition(kind=ConditionKind.Compare, left="{Rand.Valoare}",
                                       operator=CompareOperator.IsEmpty),
                   then=[SetVariableNode(name="Facut", value="completat"), WriteExcelNode(column="Nota", value="nou")],
                   else_=[SetVariableNode(name="Facut", value="exista")])]))
        out = os.path.join(d, "rapoartele_mele")
        ReportSettings(enabled=True, folder=out, variables=["Rand.ID", "Facut"]).save_to(r)
        res = RprEngine(ObjectRepository(), DriverProvider(), RunHost(), d).run(r)
        self.assertTrue(res.success, res.message)
        self.assertEqual(out, os.path.dirname(res.report_path))
        rows = _rows(openpyxl.load_workbook(res.report_path)["Rezultate"])
        head = rows[0]
        self.assertEqual(["Rand 2: A1", "Rand 3: A2", "Rand 4: A3"], [x[1] for x in rows[1:]])
        self.assertEqual(["exista", "completat", "exista"], [x[head.index("Facut")] for x in rows[1:]])
        self.assertEqual(["A1", "A2", "A3"], [x[head.index("Rand.ID")] for x in rows[1:]])
        # Excel-ul de intrare primeste in continuare statusul si nota
        src = _rows(openpyxl.load_workbook(xl).active)
        self.assertIn("Status", src[0])
        self.assertEqual("nou", src[2][src[0].index("Nota")])


if __name__ == "__main__":
    unittest.main()
